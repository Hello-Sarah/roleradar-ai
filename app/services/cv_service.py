import hashlib
import json
import logging
import re
import uuid
from collections.abc import Callable
from contextlib import suppress
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openai import OpenAI
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import Settings
from app.database.models import CVDocument, GeneratedCV, Job
from app.schemas import CVDocumentRead, CVLibraryScanRead, GeneratedCVRead

SUPPORTED_CV_SUFFIXES = {".docx", ".pdf", ".txt"}
CV_PROMPT_VERSION = "tailored-cv-v2-document-body"
CONTROLLED_SECTION_LABELS = frozenset(
    {"Experience", "Professional Experience", "Education", "Projects", "Certifications"}
)
logger = logging.getLogger(__name__)


class CVLibraryError(ValueError):
    pass


class EvidenceBackedItem(BaseModel):
    text: str = Field(min_length=2, max_length=500)
    source_quote: str = Field(min_length=2, max_length=800)
    source_document_id: int | None = None


class TailoredCVSection(BaseModel):
    title: str = Field(min_length=2, max_length=100)
    items: list[EvidenceBackedItem] = Field(min_length=1, max_length=20)


class TailoredCVContent(BaseModel):
    name: EvidenceBackedItem
    contact: EvidenceBackedItem | None = None
    headline: EvidenceBackedItem
    summary: list[EvidenceBackedItem] = Field(min_length=1, max_length=4)
    skills: list[EvidenceBackedItem] = Field(min_length=1, max_length=25)
    sections: list[TailoredCVSection] = Field(min_length=1, max_length=8)


@dataclass(frozen=True)
class ScanCounters:
    discovered: int
    added: int
    updated: int
    unchanged: int
    failed: list[str]


@dataclass(frozen=True)
class PreparedTailoredCV:
    job_id: int
    job_input_fingerprint: str
    file_name: str
    file_path: str
    output_hash: str
    source_cv_ids: list[int]
    source_cv_hashes: dict[str, str]
    source_cv_input_fingerprints: dict[str, str]
    source_evidence: list[dict[str, object]]
    model_version: str
    prompt_version: str
    generated_at: datetime


def _configured_directory(path_value: str) -> Path:
    path = Path(path_value).expanduser().resolve()
    path.mkdir(parents=True, exist_ok=True)
    if not path.is_dir():
        raise CVLibraryError(f"CV library path is not a directory: {path}")
    return path


def _extract_cv_text(path: Path) -> str:
    if path.suffix.casefold() == ".docx":
        document = Document(path)
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            parts.extend(cell.text for row in table.rows for cell in row.cells if cell.text.strip())
        text = "\n".join(parts)
    elif path.suffix.casefold() == ".pdf":
        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    else:
        text = path.read_text(encoding="utf-8")
    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 40:
        raise CVLibraryError("No usable CV text was found")
    return text


def _fingerprint(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65_536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def list_cv_documents(db: Session) -> list[CVDocument]:
    return list(
        db.scalars(
            select(CVDocument).where(CVDocument.active.is_(True)).order_by(CVDocument.file_name)
        ).all()
    )


def scan_cv_library(db: Session, settings: Settings) -> CVLibraryScanRead:
    library = _configured_directory(settings.cv_library_path)
    files = sorted(
        path
        for path in library.rglob("*")
        if path.is_file() and path.suffix.casefold() in SUPPORTED_CV_SUFFIXES
    )
    existing = {document.file_path: document for document in db.scalars(select(CVDocument))}
    seen: set[str] = set()
    added = updated = unchanged = 0
    failed: list[str] = []

    for path in files:
        resolved = str(path.resolve())
        seen.add(resolved)
        try:
            fingerprint = _fingerprint(path)
            modified_at = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC)
            document = existing.get(resolved)
            if document and document.fingerprint == fingerprint:
                document.active = True
                unchanged += 1
                continue
            extracted_text = _extract_cv_text(path)
            if document is None:
                document = CVDocument(file_path=resolved)
                db.add(document)
                added += 1
            else:
                updated += 1
            document.file_name = path.name
            document.file_type = path.suffix.casefold().lstrip(".")
            document.fingerprint = fingerprint
            document.modified_at = modified_at
            document.extracted_text = extracted_text
            document.active = True
        except Exception as exc:
            logger.warning("Failed to index CV file %s: %s", path, exc)
            failed.append(f"{path.name}: {exc}")

    for path, document in existing.items():
        if path not in seen:
            document.active = False
    db.commit()
    documents = list_cv_documents(db)
    return CVLibraryScanRead(
        library_path=str(library),
        discovered=len(files),
        added=added,
        updated=updated,
        unchanged=unchanged,
        failed=failed,
        documents=[CVDocumentRead.model_validate(document) for document in documents],
    )


def _normalized_contains(corpus: str, quote: str) -> bool:
    def normalize(value: str) -> str:
        return re.sub(r"\s+", " ", value).strip().casefold()

    return normalize(quote) in normalize(corpus)


def _validate_evidence(content: TailoredCVContent, corpus: str | list[CVDocument]) -> None:
    items = [content.name, content.headline, *content.summary, *content.skills]
    if content.contact:
        items.append(content.contact)
    items.extend(item for section in content.sections for item in section.items)
    # A quote must be wholly contained in ONE extracted body. Prompt headers, file
    # names and text spanning two documents are never evidence. Resolve a missing
    # source ID deterministically for backwards-compatible exact-quote providers.
    bodies = (
        {document.id: document.extracted_text for document in corpus}
        if isinstance(corpus, list)
        else {None: corpus}
    )
    unsupported = []
    for item in items:
        matching_ids = [
            source_id
            for source_id, body in bodies.items()
            if (item.source_document_id is None or item.source_document_id == source_id)
            and _normalized_contains(body, item.source_quote)
        ]
        if item.text != item.source_quote or not matching_ids:
            unsupported.append(item.text)
        else:
            item.source_document_id = matching_ids[0]
    if unsupported:
        raise CVLibraryError(
            "The model produced claims without verifiable CV evidence; generation was stopped"
        )
    uncontrolled_titles = [
        section.title
        for section in content.sections
        if section.title not in CONTROLLED_SECTION_LABELS
    ]
    if uncontrolled_titles:
        raise CVLibraryError("Tailored CV sections must use controlled section labels")


def _source_corpus(documents: list[CVDocument]) -> str:
    return "\n\n".join(
        f"=== SOURCE CV {document.id}: {document.file_name} ===\n{document.extracted_text}"
        for document in documents
    )


def _generate_content(
    documents: list[CVDocument], job: Job, settings: Settings
) -> TailoredCVContent:
    if not settings.openai_api_key:
        raise CVLibraryError("OPENAI_API_KEY is required to generate a tailored CV")
    corpus = _source_corpus(documents)
    client = OpenAI(api_key=settings.openai_api_key, base_url=settings.openai_base_url)
    response = client.responses.parse(
        model=settings.openai_model,
        input=[
            {
                "role": "system",
                "content": (
                    "Create an ATS-friendly tailored CV using only facts present in the source "
                    "CVs. Never invent skills, employers, dates, achievements, metrics, education, "
                    "or credentials. Every output item's text must exactly equal an exact "
                    "source_quote copied from one extracted CV body; do not rewrite it. "
                    "Include source_document_id. Filenames and SOURCE CV headers are metadata, "
                    "never factual evidence. "
                    "Section titles must be one of: Experience, Professional Experience, "
                    "Education, Projects, or "
                    "Certifications."
                ),
            },
            {
                "role": "user",
                "content": (
                    f"TARGET JOB\nCompany: {job.company}\nTitle: {job.title}\n"
                    f"Description:\n{job.description}\n\nSOURCE CVS\n{corpus}"
                ),
            },
        ],
        text_format=TailoredCVContent,
    )
    if response.output_parsed is None:
        raise CVLibraryError("The model returned no structured CV content")
    _validate_evidence(response.output_parsed, documents)
    logger.info("Generated validated CV content for job_id=%s", job.id)
    return response.output_parsed


def _safe_filename(value: str) -> str:
    return re.sub(r"[\\/:*?\"<>|\n\r]+", " ", value).strip()[:120]


def _set_font(run, size: float, bold: bool = False, color: str = "222222") -> None:
    run.font.name = "Arial"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _build_docx(content: TailoredCVContent, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    _set_font(name.add_run(content.name.text), 20, bold=True, color="102A43")
    if content.contact:
        contact = document.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(4)
        _set_font(contact.add_run(content.contact.text), 9.5, color="52606D")
    headline = document.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline.paragraph_format.space_after = Pt(8)
    _set_font(headline.add_run(content.headline.text), 11, bold=True, color="334E68")

    def add_heading(text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(3)
        _set_font(paragraph.add_run(text.upper()), 11, bold=True, color="1F4D78")

    add_heading("Professional Summary")
    summary = document.add_paragraph()
    summary.add_run(" ".join(item.text for item in content.summary))

    add_heading("Core Skills")
    skills = document.add_paragraph()
    skills.add_run(" • ".join(item.text for item in content.skills))

    for tailored_section in content.sections:
        add_heading(tailored_section.title)
        for item in tailored_section.items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.16)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.add_run(item.text)

    document.save(output)


def _cv_snapshot(document: CVDocument) -> CVDocument:
    return CVDocument(
        id=document.id,
        file_path=document.file_path,
        file_name=document.file_name,
        file_type=document.file_type,
        fingerprint=document.fingerprint,
        modified_at=document.modified_at,
        extracted_text=document.extracted_text,
        active=document.active,
    )


def _job_snapshot(job: Job) -> Job:
    return Job(
        id=job.id,
        fingerprint=job.fingerprint,
        company=job.company,
        title=job.title,
        location=job.location,
        url=job.url,
        posting_date=job.posting_date,
        description=job.description,
        source=job.source,
        status=job.status,
    )


def _tailored_cv_job_input_fingerprint(job: Job) -> str:
    payload = {
        "company": job.company,
        "title": job.title,
        "description": job.description,
    }
    digest = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
    ).hexdigest()
    return f"tailored-cv-job-input-sha256:{digest}"


def _cv_input_fingerprint(document: CVDocument) -> str:
    payload = {
        "id": document.id,
        "file_name": document.file_name,
        "extracted_text": document.extracted_text,
    }
    digest = hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
    ).hexdigest()
    return f"tailored-cv-source-input-sha256:{digest}"


def prepare_tailored_cv(
    db: Session,
    job_id: int,
    settings: Settings,
    *,
    source_cv_ids: list[int] | None = None,
    content_provider: Callable[[list[CVDocument], Job, Settings], TailoredCVContent] | None = None,
) -> PreparedTailoredCV:
    job = db.get(Job, job_id)
    if job is None:
        raise LookupError("Job not found")
    documents = list_cv_documents(db)
    if source_cv_ids is not None:
        unique_source_ids = list(dict.fromkeys(source_cv_ids))
        by_id = {document.id: document for document in documents}
        if set(unique_source_ids) != set(by_id).intersection(unique_source_ids):
            raise CVLibraryError("A selected source CV is missing or inactive")
        documents = [by_id[source_id] for source_id in unique_source_ids]
    if not documents:
        raise CVLibraryError("The CV library is empty; scan the configured folder first")
    job_snapshot = _job_snapshot(job)
    document_snapshots = [_cv_snapshot(document) for document in documents]
    # Provider access and file generation must not hold a database transaction open.
    db.rollback()
    content = (content_provider or _generate_content)(document_snapshots, job_snapshot, settings)
    # Preserve this deterministic gate at the generation boundary so validation cannot be
    # bypassed by a future provider adapter or an internal caller.
    _validate_evidence(content, document_snapshots)
    items = [content.name, content.headline, *content.summary, *content.skills]
    if content.contact:
        items.append(content.contact)
    items.extend(item for section in content.sections for item in section.items)
    output_directory = _configured_directory(settings.generated_cv_path)
    file_name = (
        f"{_safe_filename(content.name.text)}—{_safe_filename(job_snapshot.title)}-chatgpt.docx"
    )
    artifact_directory = output_directory / "artifacts" / uuid.uuid4().hex
    output = artifact_directory / file_name
    temporary_output = output_directory / f".{uuid.uuid4().hex}.docx"
    try:
        _build_docx(content, temporary_output)
        output_hash = _fingerprint(temporary_output)
        artifact_directory.mkdir(parents=True, exist_ok=False)
        temporary_output.replace(output)
        generated_at = datetime.now(UTC)
    except Exception:
        temporary_output.unlink(missing_ok=True)
        output.unlink(missing_ok=True)
        with suppress(OSError):
            artifact_directory.rmdir()
        raise
    return PreparedTailoredCV(
        job_id=job_snapshot.id,
        job_input_fingerprint=_tailored_cv_job_input_fingerprint(job_snapshot),
        file_name=file_name,
        file_path=str(output),
        output_hash=output_hash,
        source_cv_ids=[document.id for document in document_snapshots],
        source_cv_hashes={
            str(document.id): document.fingerprint for document in document_snapshots
        },
        source_cv_input_fingerprints={
            str(document.id): _cv_input_fingerprint(document) for document in document_snapshots
        },
        source_evidence=[item.model_dump() for item in items],
        model_version=settings.openai_model,
        prompt_version=CV_PROMPT_VERSION,
        generated_at=generated_at,
    )


def discard_prepared_tailored_cv(prepared: PreparedTailoredCV) -> None:
    output = Path(prepared.file_path)
    output.unlink(missing_ok=True)
    with suppress(OSError):
        output.parent.rmdir()


def persist_prepared_tailored_cv(db: Session, prepared: PreparedTailoredCV) -> GeneratedCVRead:
    job = db.get(Job, prepared.job_id)
    if job is None or _tailored_cv_job_input_fingerprint(job) != prepared.job_input_fingerprint:
        raise CVLibraryError("The target job input changed while CV generation was being prepared")
    documents = list(
        db.scalars(
            select(CVDocument).where(
                CVDocument.id.in_(prepared.source_cv_ids), CVDocument.active.is_(True)
            )
        ).all()
    )
    current_hashes = {str(document.id): document.fingerprint for document in documents}
    if current_hashes != prepared.source_cv_hashes:
        raise CVLibraryError("Source CV evidence changed while generation was being prepared")
    current_input_fingerprints = {
        str(document.id): _cv_input_fingerprint(document) for document in documents
    }
    if current_input_fingerprints != prepared.source_cv_input_fingerprints:
        raise CVLibraryError("Source CV evidence changed while generation was being prepared")
    generated = GeneratedCV(
        job_id=prepared.job_id,
        file_name=prepared.file_name,
        file_path=prepared.file_path,
        output_hash=prepared.output_hash,
        source_cv_ids=prepared.source_cv_ids,
        source_cv_hashes=prepared.source_cv_hashes,
        source_evidence=prepared.source_evidence,
        model_version=prepared.model_version,
        prompt_version=prepared.prompt_version,
        generated_at=prepared.generated_at,
    )
    db.add(generated)
    db.flush()
    logger.info(
        "Prepared tailored CV id=%s for job_id=%s using source_cv_ids=%s at %s",
        generated.id,
        prepared.job_id,
        prepared.source_cv_ids,
        prepared.file_path,
    )
    return GeneratedCVRead(
        id=generated.id,
        job_id=prepared.job_id,
        file_name=prepared.file_name,
        file_path=prepared.file_path,
        source_cv_ids=prepared.source_cv_ids,
        source_cv_hashes=prepared.source_cv_hashes,
        source_evidence=prepared.source_evidence,
        output_hash=prepared.output_hash,
        model_version=prepared.model_version,
        prompt_version=prepared.prompt_version,
        generated_at=prepared.generated_at,
    )


def generate_tailored_cv(
    db: Session,
    job_id: int,
    settings: Settings,
    *,
    source_cv_ids: list[int] | None = None,
    commit: bool = True,
    content_provider: Callable[[list[CVDocument], Job, Settings], TailoredCVContent] | None = None,
) -> GeneratedCVRead:
    prepared = prepare_tailored_cv(
        db,
        job_id,
        settings,
        source_cv_ids=source_cv_ids,
        content_provider=content_provider,
    )
    try:
        generated = persist_prepared_tailored_cv(db, prepared)
        if commit:
            db.commit()
        return generated
    except Exception:
        db.rollback()
        discard_prepared_tailored_cv(prepared)
        raise
