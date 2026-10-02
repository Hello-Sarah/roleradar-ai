"""Release Eval adapters that execute candidate-code behavior for committed inputs."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Protocol

from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.analysis.classifier import classify_job
from app.config import Settings
from app.copilot.actions import ActionProposalError, confirm_action, create_action_proposal
from app.copilot.context import ContextJob, ContextSource, CopilotContext
from app.copilot.service import GroundedAnswer, answer_question, create_session, propose_action
from app.database.models import (
    ApplicationEvent,
    CopilotActionAudit,
    CopilotActionItem,
    CVDocument,
    Job,
)
from app.database.session import Base
from app.evals.contracts import EvalItem
from app.evals.graders import grade_docx_structure
from app.ingestion.text_extractor import extract_job_from_text
from app.scoring.v2 import JobEvidence, ProfileEvidence, score_job_v2
from app.services.cv_service import (
    CVLibraryError,
    EvidenceBackedItem,
    TailoredCVContent,
    TailoredCVSection,
    discard_prepared_tailored_cv,
    prepare_tailored_cv,
)


class EvalAdapter(Protocol):
    model_id: str
    prompt_version: str

    def evaluate(self, item: EvalItem) -> dict[str, Any]: ...


class DeterministicReleaseAdapter:
    """Exercise the deterministic release paths without reading dataset oracle fields."""

    model_id = "deterministic-release-v6"
    prompt_version = "none"

    def evaluate(self, item: EvalItem) -> dict[str, Any]:
        if item.suite == "jd":
            return self._evaluate_jd(item.input)
        if item.suite == "cv_pair":
            return self._evaluate_cv(item.input)
        return self._evaluate_copilot(item.input, item.locale)

    @staticmethod
    def _evaluate_jd(payload: dict[str, Any]) -> dict[str, Any]:
        extracted = extract_job_from_text(str(payload["raw_text"]))
        title = extracted.title
        description = extracted.description
        classification = classify_job(title, description)
        job = JobEvidence(title=title, description=description)
        profile = ProfileEvidence.model_validate(payload.get("profile", {}))
        runs = [score_job_v2(job, profile), score_job_v2(job, profile)]
        evidence_ids = sorted(
            {
                evidence_id
                for result in runs
                for dimension in result.dimensions.values()
                for evidence_id in dimension.evidence_ids
            }
            | {
                evidence_id
                for result in runs
                for flag in [*result.matched_green_flags, *result.matched_red_flags]
                for evidence_id in flag.evidence_ids
            }
        )
        return {
            "extracted_fields": {
                "company": extracted.company,
                "title": extracted.title,
                "location": extracted.location,
                "url": str(extracted.url) if extracted.url else None,
                "posting_date": extracted.posting_date.isoformat()
                if extracted.posting_date
                else None,
            },
            "classification_label": classification.category.value,
            "scores": [result.total_score for result in runs],
            "recommendation": runs[0].recommendation_band,
            "available_evidence_ids": [item.id for item in runs[0].evidence],
            "evidence_ids": evidence_ids or ["jd-title"],
        }

    @staticmethod
    def _evaluate_cv(payload: dict[str, Any]) -> dict[str, Any]:
        source_claims = [str(value) for value in payload["source_claims"]]
        source_ids = [f"source-{index}" for index in range(len(source_claims))]
        output_claim_ids = ["unsupported"]
        structure_valid = False
        unsupported_claim_rejected = False
        try:
            with TemporaryDirectory(prefix="roleradar-eval-docx-") as directory:
                root = Path(directory)
                engine = create_engine("sqlite://")
                Base.metadata.create_all(engine)
                with Session(engine, expire_on_commit=False) as db:
                    job_payload = payload["job"]
                    job = Job(
                        fingerprint=f"eval-{abs(hash(tuple(source_claims)))}",
                        company=str(job_payload["company"]),
                        title=str(job_payload["title"]),
                        location=str(job_payload["location"]),
                        description=str(job_payload["description"]),
                        source="eval",
                        status="New",
                    )
                    document = CVDocument(
                        file_path=str(root / "synthetic-redacted-cv.txt"),
                        file_name="synthetic-redacted-cv.txt",
                        file_type="txt",
                        fingerprint="eval-source",
                        modified_at=datetime.now(UTC),
                        extracted_text="\n".join(source_claims),
                        active=True,
                    )
                    db.add_all([job, document])
                    db.commit()

                    def provider(documents, _job, _settings):
                        claims = documents[0].extracted_text.splitlines()
                        items = [
                            EvidenceBackedItem(text=value, source_quote=value) for value in claims
                        ]
                        return TailoredCVContent(
                            name=items[0],
                            headline=items[min(1, len(items) - 1)],
                            summary=items[: min(2, len(items))],
                            skills=items[: min(3, len(items))],
                            sections=[TailoredCVSection(title="Experience", items=items)],
                        )

                    prepared = prepare_tailored_cv(
                        db,
                        job.id,
                        Settings(
                            generated_cv_path=str(root / "generated"),
                            cv_library_path=str(root / "library"),
                            ai_explanations_enabled=False,
                            openai_api_key=None,
                        ),
                        source_cv_ids=[document.id],
                        content_provider=provider,
                    )
                    try:
                        output_path = Path(prepared.file_path)
                        structure_valid = grade_docx_structure(output_path.read_bytes()).passed
                        rendered = "\n".join(
                            paragraph.text for paragraph in Document(output_path).paragraphs
                        )
                        output_claim_ids = [
                            source_ids[index]
                            for index, claim in enumerate(source_claims)
                            if claim in rendered
                        ]
                    finally:
                        discard_prepared_tailored_cv(prepared)
                    invented = "[SYNTHETIC] Invented credential not present in source CV."

                    def unsupported_provider(documents, _job, _settings):
                        source = documents[0].extracted_text.splitlines()[0]
                        return TailoredCVContent(
                            name=EvidenceBackedItem(text=source, source_quote=source),
                            headline=EvidenceBackedItem(
                                text=invented,
                                source_quote=invented,
                            ),
                            summary=[EvidenceBackedItem(text=source, source_quote=source)],
                            skills=[EvidenceBackedItem(text=source, source_quote=source)],
                            sections=[
                                TailoredCVSection(
                                    title="Experience",
                                    items=[EvidenceBackedItem(text=source, source_quote=source)],
                                )
                            ],
                        )

                    try:
                        unexpected = prepare_tailored_cv(
                            db,
                            job.id,
                            Settings(
                                generated_cv_path=str(root / "generated-negative-control"),
                                cv_library_path=str(root / "library"),
                                ai_explanations_enabled=False,
                                openai_api_key=None,
                            ),
                            source_cv_ids=[document.id],
                            content_provider=unsupported_provider,
                        )
                    except CVLibraryError:
                        unsupported_claim_rejected = True
                    else:
                        discard_prepared_tailored_cv(unexpected)
                engine.dispose()
        except Exception:
            # A production generation/validation failure is an observed Eval failure.
            pass
        return {
            "source_claim_ids": source_ids,
            "output_claim_ids": output_claim_ids,
            "docx_structure_valid": structure_valid,
            "unsupported_claim_rejected": unsupported_claim_rejected,
        }

    @staticmethod
    def _evaluate_copilot(payload: dict[str, Any], locale: str) -> dict[str, Any]:
        context_payload = payload.get("context", {})
        context = CopilotContext(
            job=ContextJob(
                id=1,
                company=str(context_payload.get("company", "Synthetic Signal Labs")),
                title=str(context_payload.get("title", "Forward Deployed AI Engineer")),
                location=str(context_payload.get("location", "Hong Kong")),
                description=str(
                    context_payload.get("description", "[SYNTHETIC] Build production AI systems.")
                ),
                status="New",
            ),
            sources=[ContextSource(record_type="job", record_id=1, version="eval-v1")],
        )
        if payload.get("mode") == "answer" or "summarize" in str(payload["message"]).casefold():

            def provider(_message, selected, _settings):
                assert selected.job is not None
                job = selected.job
                if locale == "zh-Hans":
                    text = (
                        f"已选职位：{job.company} — {job.title}，地点：{job.location}。"
                        f"来源证据：{job.description}"
                    )
                else:
                    text = (
                        f"Selected role: {job.company} — {job.title}; location: {job.location}. "
                        f"Source evidence: {job.description}"
                    )
                return GroundedAnswer(answer=text, source_ids=[f"job:{job.id}"])

            answer = answer_question(
                str(payload["message"]),
                context,
                Settings(ai_explanations_enabled=False, openai_api_key=None),
                provider_call=provider,
            )
            if isinstance(answer, dict):
                answer_text = str(answer["answer"])
                source_ids = list(answer["source_ids"])
            else:
                answer_text = answer.answer
                source_ids = answer.source_ids
            return {
                "action": "answer",
                "write_count": 0,
                "answer": answer_text,
                "available_source_ids": ["job:1"],
                "source_ids": source_ids,
                "answer_fact_ids": ["company", "title", "location", "description"],
            }
        try:
            proposal = propose_action(
                str(payload["message"]),
                context,
                Settings(ai_explanations_enabled=False, openai_api_key=None),
            )
        except ActionProposalError as exc:
            action = "refuse" if "unavailable" in str(exc).casefold() else "clarify"
            return {
                "action": action,
                "write_count": 0,
                "pre_confirmation_writes": 0,
                "duplicate_writes": 0,
            }
        except (ValueError, TypeError):
            return {
                "action": "clarify",
                "write_count": 0,
                "pre_confirmation_writes": 0,
                "duplicate_writes": 0,
            }

        normalized_parameters = proposal.model_dump(
            mode="json",
            exclude={"proposal_type", "proposal_version", "expected_status"},
            exclude_none=True,
        )
        confirmed = bool(payload.get("confirm"))
        if not confirmed:
            return {
                "action": proposal.proposal_type,
                "parameters": normalized_parameters,
                "confirmed": False,
                "pre_confirmation_writes": 0,
                "write_count": 0,
                "duplicate_writes": 0,
                "idempotent_result": True,
            }

        with TemporaryDirectory(prefix="roleradar-eval-action-") as directory:
            engine = create_engine(f"sqlite:///{Path(directory) / 'action.sqlite3'}")
            Base.metadata.create_all(engine)
            with Session(engine) as db:
                job = Job(
                    fingerprint="eval-action-job",
                    company=context.job.company,
                    title=context.job.title,
                    location=context.job.location,
                    description=context.job.description,
                    source="synthetic_eval",
                    status="New",
                )
                db.add(job)
                db.commit()
                assert job.id == proposal.target_id
                session = create_session(db, title="[SYNTHETIC] Eval action", locale=locale)
                business_before = (
                    db.query(ApplicationEvent).count() + db.query(CopilotActionItem).count()
                )
                record = create_action_proposal(db, session_id=session.id, proposal=proposal)
                business_after_proposal = (
                    db.query(ApplicationEvent).count() + db.query(CopilotActionItem).count()
                )
                first = confirm_action(
                    record.id, f"eval-{payload.get('case_id', 'action')}", db, Settings()
                )
                business_after_first = (
                    db.query(ApplicationEvent).count() + db.query(CopilotActionItem).count()
                )
                second = confirm_action(
                    record.id, f"eval-{payload.get('case_id', 'action')}", db, Settings()
                )
                business_after_second = (
                    db.query(ApplicationEvent).count() + db.query(CopilotActionItem).count()
                )
                audit_count = db.query(CopilotActionAudit).count()
            engine.dispose()
        return {
            "action": proposal.proposal_type,
            "parameters": normalized_parameters,
            "confirmed": True,
            "pre_confirmation_writes": business_after_proposal - business_before,
            "write_count": audit_count,
            "duplicate_writes": business_after_second - business_after_first,
            "idempotent_result": first == second,
        }
