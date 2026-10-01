import hashlib
from datetime import UTC
from pathlib import Path

import pytest
from docx import Document

from app.config import Settings
from app.database.models import GeneratedCV, Job
from app.services.cv_service import (
    CVLibraryError,
    EvidenceBackedItem,
    TailoredCVContent,
    TailoredCVSection,
    _build_docx,
    _safe_filename,
    _validate_evidence,
    generate_tailored_cv,
    scan_cv_library,
)


def test_scan_cv_library_adds_and_reuses_fingerprint(db, tmp_path) -> None:
    source = tmp_path / "cv-library"
    source.mkdir()
    (source / "main.txt").write_text(
        "Jane Doe\nApplied AI Engineer\nBuilt and deployed Python APIs for banking users.",
        encoding="utf-8",
    )
    settings = Settings(cv_library_path=str(source))

    first = scan_cv_library(db, settings)
    second = scan_cv_library(db, settings)

    assert first.discovered == 1
    assert first.added == 1
    assert first.documents[0].file_name == "main.txt"
    assert second.added == 0
    assert second.unchanged == 1
    assert second.documents[0].id == first.documents[0].id


def test_scan_preserves_source_bytes_and_never_initializes_a_model(
    db, tmp_path, monkeypatch
) -> None:
    source = tmp_path / "cv-library"
    source.mkdir()
    source_file = source / "main.txt"
    source_file.write_text(
        "Jane Doe\nApplied AI Engineer\nBuilt and deployed Python APIs for banking users.",
        encoding="utf-8",
    )
    original_hash = hashlib.sha256(source_file.read_bytes()).hexdigest()
    monkeypatch.setattr(
        "app.services.cv_service.OpenAI",
        lambda *args, **kwargs: pytest.fail("Scanning must not initialize an LLM client"),
    )

    scan_cv_library(db, Settings(cv_library_path=str(source)))

    assert hashlib.sha256(source_file.read_bytes()).hexdigest() == original_hash


def test_scan_supported_formats_and_deactivates_removed_sources(db, tmp_path) -> None:
    source = tmp_path / "cv-library"
    source.mkdir()
    (source / "profile.txt").write_text(
        "Synthetic Candidate\nBuilt production AI systems with Python APIs for banking.",
        encoding="utf-8",
    )
    document = Document()
    document.add_paragraph(
        "Synthetic Candidate built production AI systems with Python APIs for banking."
    )
    document.save(source / "profile.docx")
    stream = (
        b"BT /F1 12 Tf 72 720 Td (Synthetic Candidate built production AI systems "
        b"with Python APIs for banking.) Tj ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf.extend(f"{index} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    (source / "profile.pdf").write_bytes(pdf)
    settings = Settings(cv_library_path=str(source))

    first = scan_cv_library(db, settings)
    (source / "profile.pdf").unlink()
    second = scan_cv_library(db, settings)

    assert first.discovered == 3
    assert {item.file_type for item in first.documents} == {"docx", "pdf", "txt"}
    assert second.discovered == 2
    assert {item.file_type for item in second.documents} == {"docx", "txt"}


def _validated_content() -> TailoredCVContent:
    return TailoredCVContent(
        name=EvidenceBackedItem(text="Jane Doe", source_quote="Jane Doe"),
        contact=EvidenceBackedItem(
            text="jane@example.com · Hong Kong",
            source_quote="jane@example.com · Hong Kong",
        ),
        headline=EvidenceBackedItem(text="Applied AI Engineer", source_quote="Applied AI Engineer"),
        summary=[
            EvidenceBackedItem(
                text="Built reliable Python APIs for banking users.",
                source_quote="Built reliable Python APIs for banking users.",
            )
        ],
        skills=[EvidenceBackedItem(text="Python", source_quote="Python")],
        sections=[
            TailoredCVSection(
                title="Experience",
                items=[
                    EvidenceBackedItem(
                        text="Built reliable Python APIs for banking users.",
                        source_quote="Built reliable Python APIs for banking users.",
                    )
                ],
            )
        ],
    )


def _job(db) -> Job:
    job = Job(
        fingerprint="tailored-cv-job",
        company="Example AI",
        title="Applied AI Engineer",
        location="Hong Kong",
        description="Build reliable applied AI systems for banking users.",
        source="manual",
        status="Saved",
    )
    db.add(job)
    db.commit()
    return job


def test_generation_preserves_source_bytes_and_stores_complete_provenance(
    db, tmp_path, monkeypatch
) -> None:
    source = tmp_path / "cv-library"
    output = tmp_path / "generated"
    source.mkdir()
    source_file = source / "main.txt"
    source_file.write_text(
        "Jane Doe\njane@example.com · Hong Kong\nApplied AI Engineer\n"
        "Built reliable Python APIs for banking users.",
        encoding="utf-8",
    )
    supplemental_source = source / "supplemental.txt"
    supplemental_source.write_text(
        "Python automation experience supporting reliable banking data workflows.",
        encoding="utf-8",
    )
    source_hash = hashlib.sha256(source_file.read_bytes()).hexdigest()
    supplemental_hash = hashlib.sha256(supplemental_source.read_bytes()).hexdigest()
    scan = scan_cv_library(db, Settings(cv_library_path=str(source)))
    job = _job(db)
    settings = Settings(
        cv_library_path=str(source),
        generated_cv_path=str(output),
        openai_api_key="test-key",
        openai_model="test-model",
    )
    monkeypatch.setattr(
        "app.services.cv_service._generate_content", lambda *_: _validated_content()
    )

    result = generate_tailored_cv(db, job.id, settings)

    generated = db.get(GeneratedCV, result.id)
    assert hashlib.sha256(source_file.read_bytes()).hexdigest() == source_hash
    assert result.file_name == "Jane Doe—Applied AI Engineer-chatgpt.docx"
    assert generated is not None
    assert generated.job_id == job.id
    assert generated.source_cv_ids == [document.id for document in scan.documents]
    assert generated.source_cv_hashes == {
        str(scan.documents[0].id): source_hash,
        str(scan.documents[1].id): supplemental_hash,
    }
    assert generated.model_version == "test-model"
    assert generated.prompt_version
    # SQLite does not round-trip timezone metadata, unlike the production database.
    assert generated.generated_at.replace(tzinfo=UTC) == result.generated_at
    assert generated.output_hash == hashlib.sha256(Path(result.file_path).read_bytes()).hexdigest()


def test_unsupported_generation_stops_before_writing_a_file(db, tmp_path, monkeypatch) -> None:
    source = tmp_path / "cv-library"
    output = tmp_path / "generated"
    source.mkdir()
    (source / "main.txt").write_text(
        "Jane Doe\nApplied AI Engineer\nBuilt reliable Python APIs for banking users.",
        encoding="utf-8",
    )
    scan_cv_library(db, Settings(cv_library_path=str(source)))
    job = _job(db)
    unsupported = _validated_content().model_copy(
        update={
            "summary": [
                EvidenceBackedItem(
                    text="Invented executive leadership", source_quote="not in source"
                )
            ]
        }
    )
    monkeypatch.setattr("app.services.cv_service._generate_content", lambda *_: unsupported)

    with pytest.raises(CVLibraryError, match="without verifiable CV evidence"):
        generate_tailored_cv(
            db,
            job.id,
            Settings(
                cv_library_path=str(source),
                generated_cv_path=str(output),
                openai_api_key="test-key",
            ),
        )

    assert not output.exists()
    assert db.query(GeneratedCV).count() == 0


def test_build_docx_creates_readable_tailored_cv(tmp_path) -> None:
    content = TailoredCVContent(
        name=EvidenceBackedItem(text="Jane Doe", source_quote="Jane Doe"),
        contact=EvidenceBackedItem(
            text="jane@example.com · Hong Kong",
            source_quote="jane@example.com · Hong Kong",
        ),
        headline=EvidenceBackedItem(text="Applied AI Engineer", source_quote="Applied AI Engineer"),
        summary=[
            EvidenceBackedItem(
                text="Built production AI systems for banking users.",
                source_quote="Built production AI systems for banking users.",
            )
        ],
        skills=[EvidenceBackedItem(text="Python", source_quote="Python")],
        sections=[
            TailoredCVSection(
                title="Experience",
                items=[
                    EvidenceBackedItem(
                        text="Deployed a Python API.", source_quote="Deployed a Python API."
                    )
                ],
            )
        ],
    )
    output = tmp_path / "Jane Doe—Applied AI Engineer-chatgpt.docx"

    _build_docx(content, output)

    rendered = Document(output)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "Jane Doe" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert "Deployed a Python API." in text


def test_evidence_validation_rejects_unsupported_claim() -> None:
    content = TailoredCVContent(
        name=EvidenceBackedItem(text="Jane Doe", source_quote="Jane Doe"),
        headline=EvidenceBackedItem(text="AI Engineer", source_quote="AI Engineer"),
        summary=[EvidenceBackedItem(text="Invented claim", source_quote="not in source")],
        skills=[EvidenceBackedItem(text="Python", source_quote="Python")],
        sections=[
            TailoredCVSection(
                title="Experience",
                items=[EvidenceBackedItem(text="Built APIs", source_quote="Built APIs")],
            )
        ],
    )

    try:
        _validate_evidence(content, "Jane Doe AI Engineer Python Built APIs")
    except ValueError as exc:
        assert "without verifiable CV evidence" in str(exc)
    else:
        raise AssertionError("Unsupported claims must be rejected")


def test_evidence_validation_rejects_claim_text_that_does_not_match_its_quote() -> None:
    content = _validated_content().model_copy(
        update={"skills": [EvidenceBackedItem(text="Led 100 engineers", source_quote="Python")]}
    )

    with pytest.raises(CVLibraryError, match="without verifiable CV evidence"):
        _validate_evidence(
            content,
            "Jane Doe jane@example.com · Hong Kong Applied AI Engineer Python "
            "Built reliable Python APIs for banking users.",
        )


def test_evidence_validation_rejects_invented_section_title() -> None:
    content = _validated_content().model_copy(
        update={
            "sections": [
                TailoredCVSection(
                    title="Invented Employer — 2019 to 2024",
                    items=[
                        EvidenceBackedItem(
                            text="Built reliable Python APIs for banking users.",
                            source_quote="Built reliable Python APIs for banking users.",
                        )
                    ],
                )
            ]
        }
    )

    with pytest.raises(CVLibraryError, match="controlled section labels"):
        _validate_evidence(
            content,
            "Jane Doe jane@example.com · Hong Kong Applied AI Engineer Python "
            "Built reliable Python APIs for banking users.",
        )


def test_commit_failure_leaves_no_published_artifact_or_provenance(
    db, tmp_path, monkeypatch
) -> None:
    source = tmp_path / "cv-library"
    output = tmp_path / "generated"
    source.mkdir()
    (source / "main.txt").write_text(
        "Jane Doe\njane@example.com · Hong Kong\nApplied AI Engineer\n"
        "Built reliable Python APIs for banking users.",
        encoding="utf-8",
    )
    scan_cv_library(db, Settings(cv_library_path=str(source)))
    job = _job(db)
    monkeypatch.setattr(
        "app.services.cv_service._generate_content", lambda *_: _validated_content()
    )

    def fail_commit() -> None:
        raise RuntimeError("database unavailable")

    monkeypatch.setattr(db, "commit", fail_commit)
    with pytest.raises(RuntimeError, match="database unavailable"):
        generate_tailored_cv(
            db,
            job.id,
            Settings(
                cv_library_path=str(source),
                generated_cv_path=str(output),
                openai_api_key="test-key",
            ),
        )

    assert not list(output.rglob("*.docx"))
    assert db.query(GeneratedCV).count() == 0


def test_post_commit_refresh_failure_cannot_delete_durable_artifact(
    db, tmp_path, monkeypatch
) -> None:
    source = tmp_path / "cv-library"
    output = tmp_path / "generated"
    source.mkdir()
    (source / "main.txt").write_text(
        "Jane Doe\njane@example.com · Hong Kong\nApplied AI Engineer\n"
        "Built reliable Python APIs for banking users.",
        encoding="utf-8",
    )
    scan_cv_library(db, Settings(cv_library_path=str(source)))
    job = _job(db)
    monkeypatch.setattr(
        "app.services.cv_service._generate_content", lambda *_: _validated_content()
    )
    monkeypatch.setattr(
        db,
        "refresh",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("refresh unavailable")),
    )

    generated = generate_tailored_cv(
        db,
        job.id,
        Settings(
            cv_library_path=str(source),
            generated_cv_path=str(output),
            openai_api_key="test-key",
        ),
    )

    stored = db.get(GeneratedCV, generated.id)
    assert stored is not None
    assert Path(stored.file_path).is_file()
    assert hashlib.sha256(Path(stored.file_path).read_bytes()).hexdigest() == stored.output_hash


def test_repeated_generation_preserves_each_artifact_bytes_and_provenance(
    db, tmp_path, monkeypatch
) -> None:
    source = tmp_path / "cv-library"
    output = tmp_path / "generated"
    source.mkdir()
    (source / "main.txt").write_text(
        "Jane Doe\njane@example.com · Hong Kong\nApplied AI Engineer\n"
        "Built reliable Python APIs for banking users.",
        encoding="utf-8",
    )
    scan_cv_library(db, Settings(cv_library_path=str(source)))
    job = _job(db)
    settings = Settings(
        cv_library_path=str(source),
        generated_cv_path=str(output),
        openai_api_key="test-key",
    )
    monkeypatch.setattr(
        "app.services.cv_service._generate_content", lambda *_: _validated_content()
    )

    first = generate_tailored_cv(db, job.id, settings)
    first_bytes = Path(first.file_path).read_bytes()
    second = generate_tailored_cv(db, job.id, settings)

    assert first.file_name == second.file_name == "Jane Doe—Applied AI Engineer-chatgpt.docx"
    assert first.file_path != second.file_path
    assert Path(first.file_path).read_bytes() == first_bytes
    assert hashlib.sha256(first_bytes).hexdigest() == first.output_hash
    assert hashlib.sha256(Path(second.file_path).read_bytes()).hexdigest() == second.output_hash
    assert db.query(GeneratedCV).count() == 2


def test_safe_filename_removes_path_characters() -> None:
    assert _safe_filename("Jane/Smith: AI*Lead?") == "Jane Smith  AI Lead"
